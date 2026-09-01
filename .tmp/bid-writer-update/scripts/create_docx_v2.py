#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Create a bid DOCX from Markdown with section-mapped evidence images.

This is a reusable bid-writer template. It intentionally does not create a
generic appendix or attachment dump for company material images. Images are
embedded only when they can be mapped to a matching body section; unmapped
images are reported for manual placement decisions.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
DEFAULT_CATEGORY_SECTIONS = {
    "business_license": ["资信", "商务", "资格", "资质", "营业执照"],
    "qualification": ["资信", "商务", "资格", "资质", "认证", "证书"],
    "software_copyright": ["软件著作权", "著作权", "资质", "知识产权"],
    "patent": ["专利", "知识产权", "资质"],
    "contract_case": ["业绩", "案例", "合同", "类似项目"],
    "personnel": ["人员", "项目团队", "项目实施人员", "人员一览表"],
    "tax_audit": ["财务", "纳税", "审计", "社保"],
    "authorization": ["授权", "原厂", "制造商", "售后服务"],
    "hardware_product": ["货物性能", "配置", "设备", "产品", "技术响应"],
    "product_sheet": ["产品", "彩页", "技术响应", "货物性能", "配置"],
    "screenshot": ["截图", "界面", "系统", "平台", "技术方案"],
    "seal_signature": ["签字", "盖章", "授权委托", "法定代表人"],
}
APPENDIX_DUMP_HEADING_RE = re.compile(
    r"(?:\u9644\u5f55|\u9644\u4ef6|\u9644\u56fe)\s*[:\uff1a]?\s*"
    r"(?:\u4f01\u4e1a\u8d44\u6599\u56fe\u7247|\u4f01\u4e1a\u8d44\u6599|"
    r"\u8bc1\u660e\u6750\u6599|\u8bc1\u636e\u56fe\u7247|\u56fe\u7247)|"
    r"\u4f01\u4e1a\u8d44\u6599\u56fe\u7247|"
    r"appendix\s*[:\uff1a]?\s*(?:company|enterprise|material).*image|"
    r"remaining\s+images|all\s+remaining\s+images",
    re.I,
)
GENERIC_IMAGE_TEXT_RE = re.compile(
    r"\u8bc1\u636e\u56fe\u7247\s*[:\uff1a]|"
    r"\u672c\u9875\u4e3a\u4e0e\u672c\u8282\u5185\u5bb9\u5bf9\u5e94\u7684"
    r"\u4f01\u4e1a\u8d44\u6599\u8bc1\u660e\u56fe\u7247|"
    r"\u4f01\u4e1a\u8d44\u6599\u8bc1\u660e\u56fe\u7247",
    re.I,
)


@dataclass
class ImageItem:
    path: Path
    relative_path: str
    name: str
    categories: list[str] = field(default_factory=list)


@dataclass
class AssetRule:
    section: str
    patterns: list[str] = field(default_factory=list)
    categories: list[str] = field(default_factory=list)
    max_images: int | None = None


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value).lower()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def load_inventory(path: Path | None) -> list[ImageItem]:
    if not path:
        return []
    data = read_json(path)
    images: list[ImageItem] = []
    for item in data.get("files", []):
        if not item.get("is_image"):
            continue
        raw_path = item.get("path")
        if not raw_path:
            continue
        images.append(
            ImageItem(
                path=Path(raw_path),
                relative_path=str(item.get("relative_path") or item.get("name") or raw_path),
                name=str(item.get("name") or Path(raw_path).name),
                categories=[str(category) for category in item.get("categories", [])],
            )
        )
    return images


def iter_image_files(root: Path | None) -> Iterable[ImageItem]:
    if not root:
        return []
    if root.is_file() and root.suffix.lower() in IMAGE_EXTS:
        return [ImageItem(path=root, relative_path=root.name, name=root.name)]
    if not root.exists():
        return []
    out: list[ImageItem] = []
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in IMAGE_EXTS:
            try:
                relative = str(path.relative_to(root))
            except ValueError:
                relative = path.name
            out.append(ImageItem(path=path, relative_path=relative, name=path.name))
    return out


def load_asset_rules(path: Path | None) -> list[AssetRule]:
    if not path:
        return []
    data = read_json(path)
    rules: list[AssetRule] = []
    raw_rules = data.get("mappings", data if isinstance(data, list) else [])
    if not isinstance(raw_rules, list):
        raise SystemExit("asset map must be a list or contain a 'mappings' list")
    for raw in raw_rules:
        if not isinstance(raw, dict):
            continue
        section = str(raw.get("section") or "").strip()
        if not section:
            continue
        max_images = raw.get("max_images")
        rules.append(
            AssetRule(
                section=section,
                patterns=[str(value) for value in raw.get("patterns", [])],
                categories=[str(value) for value in raw.get("categories", [])],
                max_images=int(max_images) if isinstance(max_images, int) and max_images > 0 else None,
            )
        )
    return rules


def default_rules_from_inventory(images: list[ImageItem]) -> list[AssetRule]:
    sections: dict[str, AssetRule] = {}
    for image in images:
        for category in image.categories:
            section_keywords = DEFAULT_CATEGORY_SECTIONS.get(category)
            if not section_keywords:
                continue
            section = section_keywords[0]
            rule = sections.setdefault(section, AssetRule(section=section))
            rule.categories.append(category)
            rule.patterns.extend(section_keywords)
    return list(sections.values())


def image_matches_rule(image: ImageItem, rule: AssetRule) -> bool:
    haystack = normalize_text(" ".join([image.relative_path, image.name, *image.categories]))
    if rule.categories and any(category in image.categories for category in rule.categories):
        return True
    return any(normalize_text(pattern) in haystack for pattern in rule.patterns if pattern.strip())


def section_matches_rule(section_path: list[str], rule: AssetRule) -> bool:
    section_text = normalize_text(" ".join(section_path))
    return normalize_text(rule.section) in section_text or any(
        normalize_text(pattern) in section_text for pattern in [rule.section, *rule.patterns] if pattern.strip()
    )


def add_toc(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    for kind, text in [
        ("begin", None),
        ("instrText", 'TOC \\o "1-3" \\h \\z \\u'),
        ("separate", None),
        ("end", None),
    ]:
        if kind == "instrText":
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = text
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        run._r.append(node)


def set_fonts(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for section in document.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)


def add_heading(document: Document, text: str, level: int) -> None:
    level = max(1, min(level, 4))
    paragraph = document.add_heading(text, level=level)
    if level <= 2:
        paragraph.paragraph_format.page_break_before = True
    for run in paragraph.runs:
        run.font.name = "SimHei"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    run.font.name = "SimSun"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for index in range(width):
            cells[index].text = row[index] if index < len(row) else ""


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(row: list[str]) -> bool:
    return bool(row) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row)


def clean_image_title(image: ImageItem) -> str:
    title = Path(image.name).stem
    title = re.sub(r"[_\-]+", " ", title)
    title = re.sub(r"\s+", " ", title).strip()
    title = re.sub(r"(?:\s+|_)?\d+$", "", title).strip()
    return title or Path(image.relative_path).stem or image.name


def image_title_for_section(section_path: list[str], image: ImageItem) -> str:
    title = clean_image_title(image)
    if not section_path:
        return title
    section_leaf = section_path[-1].strip()
    normalized_title = normalize_text(title)
    normalized_section = normalize_text(section_leaf)
    if normalized_title and normalized_title in normalized_section:
        return section_leaf
    return title


def image_caption_for_section(section_path: list[str], image: ImageItem) -> str:
    title = image_title_for_section(section_path, image)
    if not section_path:
        return f"\u56fe\uff1a{title}"
    section_leaf = section_path[-1].strip()
    if normalize_text(title) in normalize_text(section_leaf):
        return f"\u56fe\uff1a{title}"
    return f"\u56fe\uff1a{section_leaf} - {title}"


def add_image_block(document: Document, image: ImageItem, section_path: list[str], max_width: Cm = Cm(15.5)) -> bool:
    if not image.path.exists():
        return False
    heading = document.add_paragraph()
    heading.style = document.styles["Heading 4"]
    heading.paragraph_format.page_break_before = True
    heading.add_run(image_title_for_section(section_path, image))

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image.path), width=max_width)

    caption = document.add_paragraph(image_caption_for_section(section_path, image))
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.font.size = Pt(9)
    return True


def insert_mapped_images(
    document: Document,
    section_path: list[str],
    images: list[ImageItem],
    rules: list[AssetRule],
    inserted: set[str],
) -> None:
    for rule in rules:
        if not section_matches_rule(section_path, rule):
            continue
        count = 0
        for image in images:
            key = str(image.path)
            if key in inserted:
                continue
            if not image_matches_rule(image, rule):
                continue
            if add_image_block(document, image, section_path):
                inserted.add(key)
                count += 1
            if rule.max_images is not None and count >= rule.max_images:
                break


def render_markdown(document: Document, markdown: str, images: list[ImageItem], rules: list[AssetRule]) -> set[str]:
    inserted: set[str] = set()
    section_path: list[str] = []
    lines = markdown.splitlines()
    index = 0
    skip_appendix_level: int | None = None
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            text = heading.group(2).strip()
            if skip_appendix_level is not None and level > skip_appendix_level:
                index += 1
                continue
            skip_appendix_level = None
            if APPENDIX_DUMP_HEADING_RE.search(text):
                skip_appendix_level = level
                index += 1
                continue
            if GENERIC_IMAGE_TEXT_RE.search(text):
                index += 1
                continue
            section_path = section_path[: level - 1] + [text]
            add_heading(document, text, level)
            insert_mapped_images(document, section_path, images, rules, inserted)
            index += 1
            continue
        if skip_appendix_level is not None:
            index += 1
            continue
        if GENERIC_IMAGE_TEXT_RE.search(stripped):
            index += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            rows: list[list[str]] = []
            while index < len(lines):
                current = lines[index].strip()
                if not (current.startswith("|") and current.endswith("|")):
                    break
                row = parse_table_row(current)
                if not is_table_separator(row):
                    rows.append(row)
                index += 1
            add_markdown_table(document, rows)
            continue
        add_paragraph(document, re.sub(r"\*\*(.*?)\*\*", r"\1", stripped))
        index += 1
    return inserted


def write_unmapped_report(path: Path, images: list[ImageItem], inserted: set[str]) -> None:
    unmapped = [
        {
            "path": str(image.path),
            "relative_path": image.relative_path,
            "name": image.name,
            "categories": image.categories,
            "reason": "No matching body section/rule; not appended to a generic appendix.",
        }
        for image in images
        if str(image.path) not in inserted
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            {
                "schema": "bid-writer.unmapped-images.v1",
                "total_images": len(images),
                "embedded_images": len(inserted),
                "unmapped_images": unmapped,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create a DOCX bid document with section-mapped evidence images.")
    parser.add_argument("--markdown", required=True, help="Source bid Markdown file.")
    parser.add_argument("--out", required=True, help="Output DOCX path.")
    parser.add_argument("--materials-json", help="materials_inventory.json from inventory_materials.py.")
    parser.add_argument("--materials-root", help="Fallback image root when no inventory JSON is available.")
    parser.add_argument("--asset-map", help="Optional JSON mappings: [{section, patterns, categories, max_images}].")
    parser.add_argument("--unmapped-report", default="unmapped_images.json", help="JSON report for images not embedded.")
    parser.add_argument("--title", help="Optional title paragraph before TOC.")
    parser.add_argument("--bidder", help="Optional bidder paragraph before TOC.")
    return parser.parse_args()


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    args = parse_args()
    markdown_path = Path(args.markdown)
    out_path = Path(args.out)
    images = load_inventory(Path(args.materials_json)) if args.materials_json else list(iter_image_files(Path(args.materials_root)) if args.materials_root else [])
    rules = load_asset_rules(Path(args.asset_map)) if args.asset_map else default_rules_from_inventory(images)

    document = Document()
    set_fonts(document)
    if args.title:
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(args.title)
        run.bold = True
        run.font.size = Pt(22)
        run.font.name = "SimHei"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    if args.bidder:
        bidder = document.add_paragraph()
        bidder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bidder.add_run(args.bidder)
    add_toc(document)

    markdown = markdown_path.read_text(encoding="utf-8")
    inserted = render_markdown(document, markdown, images, rules)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(out_path)
    write_unmapped_report(Path(args.unmapped_report), images, inserted)

    print(f"Generated: {out_path}")
    print(f"Embedded mapped images: {len(inserted)} / {len(images)}")
    if len(inserted) < len(images):
        print(f"Unmapped images report: {args.unmapped_report}")
        print("Unmapped images were not moved to an appendix; map them to body sections or report a draft gap.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
